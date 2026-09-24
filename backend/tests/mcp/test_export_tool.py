"""cmc_export_workbook: a signed, short-lived link and a summary, never the bytes."""

import re
from datetime import timedelta
from io import BytesIO
from urllib.parse import urlsplit

from openpyxl import load_workbook

from app.services import export_links, signed_links
from app.services.items import ItemFilters
from app.services.tokens import revoke_token
from tests.mcp.factories import make_item

LINK = re.compile(r"https?://\S+/api/export/link/(dl_\S+)")


def _link_path(text: str) -> str:
    match = LINK.search(text)
    assert match, text
    return urlsplit(match.group(0)).path


def _rows(content: bytes) -> list[tuple]:
    sheet = load_workbook(BytesIO(content)).active
    return list(sheet.iter_rows(min_row=2, values_only=True))


def test_export_returns_a_link_and_summary_not_bytes(mcp_client, cli_db, program, admin, vocab):
    make_item(cli_db, program, admin, title="Stability protocol review")
    make_item(cli_db, program, admin, title="Shipping validation", status="blocked")

    text = mcp_client.text("cmc_export_workbook")

    assert "2 items" in text
    assert "Entry No." in text and "Status Updates" in text
    assert "no filters" in text
    assert len(text) < 2000
    assert "valid until" in text


def test_export_link_downloads_the_filtered_workbook(mcp_client, cli_db, program, admin, vocab):
    make_item(cli_db, program, admin, title="Stability protocol review")
    make_item(cli_db, program, admin, title="Shipping validation", status="blocked")

    text = mcp_client.text("cmc_export_workbook", status=["blocked"])
    assert "1 item," in text
    assert "status blocked" in text

    response = mcp_client._client.get(_link_path(text))
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    titles = [row[3] for row in _rows(response.content)]
    assert titles == ["Shipping validation"]


def test_export_link_needs_no_credential(mcp_client, cli_db, program, admin, vocab):
    make_item(cli_db, program, admin)
    path = _link_path(mcp_client.text("cmc_export_workbook"))
    # The MCP client sends its bearer per request and holds no cookie: this GET
    # carries no credential at all, as a person's browser would.
    assert not mcp_client._client.cookies
    assert mcp_client._client.get(path).status_code == 200


def test_export_validates_filters_like_search(mcp_client, vocab):
    text = mcp_client.error("cmc_export_workbook", status=["finished"])
    assert "finished" in text
    assert "in_progress" in text


def test_export_with_no_matches_mints_no_link(mcp_client, cli_db, program, admin, vocab):
    make_item(cli_db, program, admin, status="open")
    text = mcp_client.text("cmc_export_workbook", status=["cancelled"])
    assert "nothing to export" in text
    assert "/api/export/link/" not in text


def test_period_report_is_refused_until_v2_builds_it(mcp_client, vocab):
    text = mcp_client.error("cmc_export_workbook", kind="period_report", from_date="2026-07-01")
    assert "period_report" in text
    assert "not available yet" in text


def test_dates_are_only_for_the_period_report(mcp_client, vocab):
    text = mcp_client.error("cmc_export_workbook", from_date="2026-07-01")
    assert "due_after" in text


def test_tampered_link_is_not_found(mcp_client, cli_db, program, admin, vocab):
    make_item(cli_db, program, admin)
    path = _link_path(mcp_client.text("cmc_export_workbook"))
    tampered = path[:-1] + ("0" if path[-1] != "0" else "1")
    assert mcp_client._client.get(tampered).status_code == 404


def test_expired_link_says_so(mcp_client, cli_db, program, admin, vocab, monkeypatch):
    make_item(cli_db, program, admin)
    monkeypatch.setattr(export_links, "LINK_TTL", timedelta(seconds=-1))
    path = _link_path(mcp_client.text("cmc_export_workbook"))
    response = mcp_client._client.get(path)
    assert response.status_code == 410
    assert "expired" in response.json()["error"]["message"]


def test_revoking_the_token_kills_its_links(mcp_client, cli_db, program, admin, vocab):
    from app.models import ApiToken

    make_item(cli_db, program, admin)
    path = _link_path(mcp_client.text("cmc_export_workbook"))
    token = cli_db.query(ApiToken).filter_by(name="Test agent").one()
    revoke_token(cli_db, actor=admin, token=token)

    response = mcp_client._client.get(path)
    assert response.status_code == 403
    assert "revoked" in response.json()["error"]["message"]


def test_a_link_signed_for_another_purpose_is_refused(client, db, admin, program):
    payload = {"u": admin.id, "t": None, "f": {}}
    token, _ = signed_links.sign("something.else", payload, timedelta(minutes=5))
    assert client.get(f"/api/export/link/{token}").status_code == 404


def test_filters_survive_the_round_trip():
    from datetime import date

    filters = ItemFilters(
        status=("open", "blocked"),
        group=("General Issues",),
        due_before=date(2026, 12, 31),
        q="stability",
    )
    assert export_links._filters_from(export_links._filters_payload(filters)) == filters
