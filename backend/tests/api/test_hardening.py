"""Regressions for the PR #6 review findings that live outside the MCP adapter."""

import zipfile
from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.constants import MAX_DB_INT, UPDATE_BODY_MAX
from app.models import ActionItem, AuditEvent, ItemUpdate
from app.models.base import utcnow
from app.schemas.items import ItemCreate, ItemPatch
from app.services import items as items_service
from app.services.errors import ConflictError
from app.services.items import create_item, patch_item
from app.services.revert import revert_event
from app.services.tokens import create_token, resolve_token

NEW = ItemCreate(title="Hardening target", group="General Issues", owner_org="gensci")
FORMULAS = ('=HYPERLINK("http://evil.example","click")', "+1+1", "-2+3", "@SUM(1)")


# --- 2. formula injection in the Excel export --------------------------------------


def test_export_writes_formula_like_text_as_text(admin_client, db, admin, program, vocab):
    item = create_item(
        db,
        actor=admin,
        program=program,
        data=ItemCreate(
            title=FORMULAS[0], group="General Issues", owner_org="gensci", notes_risks=FORMULAS[1]
        ),
    )
    patch_item(db, actor=admin, item=item, patch=ItemPatch(file_path=FORMULAS[3]))

    response = admin_client.get("/api/export/excel")
    assert response.status_code == 200
    sheet_xml = zipfile.ZipFile(BytesIO(response.content)).read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet_xml  # no cell is a formula

    row = next(load_workbook(BytesIO(response.content)).active.iter_rows(min_row=2))
    by_header = {cell.column: cell for cell in row}
    assert by_header[4].value == FORMULAS[0]  # the text survives exactly, for re-import
    assert by_header[4].data_type == "s"


# --- 3. two creates racing for the same entry number --------------------------------


def test_a_create_that_loses_the_entry_number_race_retries(db, admin, program, vocab, monkeypatch):
    first = create_item(db, actor=admin, program=program, data=NEW)
    real = items_service.next_entry_no
    stale = iter([first.entry_no])  # what a concurrent request read a moment too early

    monkeypatch.setattr(
        items_service, "next_entry_no", lambda *args: next(stale, None) or real(*args)
    )
    second = create_item(
        db, actor=admin, program=program, data=NEW.model_copy(update={"title": "Second"})
    )
    assert second.entry_no == first.entry_no + 1
    assert db.query(ActionItem).count() == 2


def test_a_create_that_keeps_losing_says_so_readably(db, admin, program, vocab, monkeypatch):
    first = create_item(db, actor=admin, program=program, data=NEW)
    monkeypatch.setattr(items_service, "next_entry_no", lambda *args: first.entry_no)
    with pytest.raises(ConflictError, match="entry number"):
        create_item(db, actor=admin, program=program, data=NEW)
    assert db.query(ActionItem).count() == 1


def test_a_racing_idempotent_create_names_the_original(db, admin, program, vocab, monkeypatch):
    first = create_item(db, actor=admin, program=program, data=NEW, idempotency_key="k-1")
    # The caller checked for the key before the other request committed it.
    with pytest.raises(ConflictError, match=f"#{first.entry_no}"):
        create_item(db, actor=admin, program=program, data=NEW, idempotency_key="k-1")


# --- 4. two people undoing the same change at once ---------------------------------


def test_only_one_of_two_concurrent_undos_lands(engine, db, admin, program, vocab):
    item = create_item(db, actor=admin, program=program, data=NEW)
    patch_item(db, actor=admin, item=item, patch=ItemPatch(priority="p1"))
    event_id = db.query(AuditEvent).filter_by(action="updated").one().id

    other: Session = Session(bind=engine, expire_on_commit=False)
    try:
        # The second request read the event and the item before the first undo landed.
        stale_event = other.get(AuditEvent, event_id)
        stale_item = other.get(ActionItem, item.id)  # held, as a live request holds it
        assert stale_item.priority == "p1"
        revert_event(db, actor=admin, event=db.get(AuditEvent, event_id))
        with pytest.raises(ConflictError, match="already undone"):
            revert_event(other, actor=admin, event=stale_event)
    finally:
        other.close()

    db.expire_all()
    assert db.query(AuditEvent).filter_by(action="reverted").count() == 1


# --- 5. ids too large for the database ----------------------------------------------


@pytest.mark.parametrize(
    "path",
    [
        f"/api/items/{2**63}",
        f"/api/items/{MAX_DB_INT + 1}/history",
        f"/api/items/{2**70}/updates",
        f"/api/items?assignee_id={2**63}",
        f"/api/activity?actor_id={2**63}",
    ],
)
def test_huge_ids_are_a_validation_error_not_a_crash(admin_client, program, vocab, path):
    response = admin_client.get(path)
    assert response.status_code == 422, response.text


def test_a_huge_id_in_a_body_is_a_validation_error(admin_client, program, vocab):
    response = admin_client.post(
        "/api/items",
        json={"title": "x", "group": "General Issues", "owner_org": "gensci", "assignee_id": 2**63},
    )
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "validation_error"


# --- 8. update bodies are capped ----------------------------------------------------


def test_an_oversized_update_body_is_refused(admin_client, db, admin, program, vocab):
    item = create_item(db, actor=admin, program=program, data=NEW)
    url = f"/api/items/{item.id}/updates"
    assert admin_client.post(url, json={"body": "x" * UPDATE_BODY_MAX}).status_code == 201
    response = admin_client.post(url, json={"body": "x" * (UPDATE_BODY_MAX + 1)})
    assert response.status_code == 422
    assert db.query(ItemUpdate).count() == 1


# --- 10. last_used_at is not written on every request -------------------------------


def test_last_used_is_written_at_most_once_a_minute(db, admin):
    token, raw = create_token(db, actor=admin, owner=admin, name="Busy agent")
    assert resolve_token(db, raw).last_used_at is not None
    first = token.last_used_at
    resolve_token(db, raw)
    assert token.last_used_at == first  # throttled

    token.last_used_at = utcnow() - timedelta(minutes=2)
    db.commit()
    assert resolve_token(db, raw).last_used_at > first


# --- 11. confirm tokens are signed with a derived key, not the raw secret -----------


def test_confirm_tokens_do_not_sign_with_the_raw_secret():
    import hashlib
    import hmac

    from app.config import get_settings
    from app.mcp import confirm

    raw = hmac.new(get_settings().secret_key.encode(), b"x", hashlib.sha256).hexdigest()
    assert hmac.new(confirm._key(), b"x", hashlib.sha256).hexdigest() != raw


# --- 9. security headers ------------------------------------------------------------


@pytest.mark.parametrize("path", ["/api/health", "/mcp/", "/api/items"])
def test_responses_carry_security_headers(client, path):
    headers = client.get(path).headers
    assert headers["x-content-type-options"] == "nosniff"
    assert headers["x-frame-options"] == "DENY"
    assert headers["referrer-policy"] == "no-referrer"
    assert "script-src 'self'" in headers["content-security-policy"]
    assert "frame-ancestors 'none'" in headers["content-security-policy"]


def test_the_api_docs_keep_working_without_the_csp(client):
    response = client.get("/api/docs")
    assert response.status_code == 200
    assert "content-security-policy" not in response.headers
    assert response.headers["x-content-type-options"] == "nosniff"
