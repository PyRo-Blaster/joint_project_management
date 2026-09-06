"""Read the 'Action Item' sheet into raw rows keyed by canonical column names."""

from dataclasses import dataclass
from pathlib import Path
from typing import IO, Any

import openpyxl

from app.services.errors import ImportFormatError

SHEET_NAME = "Action Item"
EXPECTED_HEADERS = {
    "entry no.": "entry_no",
    "date": "date",
    "group": "group",
    "action item": "action_item",
    "translation in mandarin": "translation",
    "owner": "owner",
    "cmc category": "category",
    "status": "status",
    "checkpoint/ddl": "due",
    "priority (p1 as highest)": "priority",
    "status updates": "status_updates",
    "notes/risks": "notes_risks",
    "file path": "file_path",
}


@dataclass(frozen=True)
class RawRow:
    excel_row: int
    values: dict[str, Any]


def _header_key(cell: Any) -> str:
    return str(cell).strip().rstrip(":：").strip().lower()


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _column_map(header: tuple) -> dict[int, str]:
    mapping = {
        index: EXPECTED_HEADERS[_header_key(cell)]
        for index, cell in enumerate(header)
        if cell is not None and _header_key(cell) in EXPECTED_HEADERS
    }
    missing = set(EXPECTED_HEADERS.values()) - set(mapping.values())
    if missing:
        raise ImportFormatError(f"Missing columns: {', '.join(sorted(missing))}")
    return mapping


def read_rows(source: str | Path | IO[bytes]) -> list[RawRow]:
    try:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for unreadable files
        raise ImportFormatError(f"Could not open workbook: {exc}") from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ImportFormatError(
                f"Sheet '{SHEET_NAME}' not found; sheets present: {workbook.sheetnames}"
            )
        rows_iter = workbook[SHEET_NAME].iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise ImportFormatError("Sheet is empty")
        mapping = _column_map(header)
        rows = []
        for excel_row, values in enumerate(rows_iter, start=2):
            record = {
                name: (values[index] if index < len(values) else None)
                for index, name in mapping.items()
            }
            if all(_is_empty(value) for value in record.values()):
                continue
            rows.append(RawRow(excel_row=excel_row, values=record))
        return rows
    finally:
        workbook.close()
